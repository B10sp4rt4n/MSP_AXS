"""Import safety: permissions, no preview writes, atomicity and replay."""
import base64
import io
import re
from zipfile import ZipFile
from pathlib import Path
import pytest
from openpyxl import load_workbook
from tests.test_personal_condominio import people, headers
from backend.db.core import Casa, Usuario, EventOutbox
from backend.services.resident_import import parse_excel
from fastapi import HTTPException

TEMPLATE = Path('backend/static/plantilla_residentes.xlsx')
URL = '/condominios/a1/residentes/importar'


def without_dimensions(raw):
    out = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as source, ZipFile(out, 'w') as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename.startswith('xl/worksheets/'):
                data = re.sub(rb'<(?:\w+:)?dimension\b[^>]*/>', b'', data)
            target.writestr(entry, data)
    return out.getvalue()


def test_shipped_template_declares_dimensions():
    book = load_workbook(TEMPLATE, read_only=True)
    try:
        assert book['Residentes'].max_column == 4
        assert book['Residentes'].max_row == 501
    finally:
        book.close()


def test_missing_dimensions_preserves_values():
    raw = base64.b64decode(body([['001','casa','Ana','ANA@example.com']])['archivo'])
    rows = parse_excel(without_dimensions(raw))
    assert len(rows) == 1
    assert rows[0]['vivienda'] == '001'
    assert rows[0]['correo'] == 'ana@example.com'
    assert rows[0]['errores'] == []


@pytest.mark.parametrize('omit_dimensions', [False, True])
@pytest.mark.parametrize('extra', ['rows', 'columns'])
def test_dimensions_limits_still_apply(omit_dimensions, extra):
    rows = [['001','casa','Ana','ana@example.com']]
    if extra == 'rows': rows *= 501
    else: rows[0].append('unexpected')
    raw = base64.b64decode(body(rows)['archivo'])
    if omit_dimensions: raw = without_dimensions(raw)
    with pytest.raises(HTTPException, match='') as exc:
        parse_excel(raw)
    assert exc.value.status_code == 400
    assert 'cuatro columnas' in exc.value.detail


def test_missing_dimensions_empty_template():
    with pytest.raises(HTTPException) as exc:
        parse_excel(without_dimensions(TEMPLATE.read_bytes()))
    assert 'no contiene residentes' in exc.value.detail


def body(rows):
    # Test fixtures use the shipped template, including its text formatting.
    book = load_workbook(TEMPLATE)
    sheet = book['Residentes']
    for n, row in enumerate(rows, 2):
        for col, value in enumerate(row, 1): sheet.cell(n, col, value)
    raw = io.BytesIO(); book.save(raw); book.close()
    return {'archivo': base64.b64encode(raw.getvalue()).decode()}


def test_preview_commit_replay_and_audit(client, people):
    auth = headers('admin', 'ADMIN_CONDOMINIO')
    payload = body([['A-001', 'depto', 'Ana', 'ANA@example.com']])
    response = client.post(URL, headers=auth, json=payload)
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview['nuevos'] == 1 and preview['errores'] == 0
    assert people.query(Casa).count() == 0
    result = client.post(URL, headers=auth, json={**payload, 'revision': preview['revision']})
    assert result.status_code == 200, result.text
    assert result.json()['creados'] == 1
    user = people.query(Usuario).filter_by(email='ana@example.com').one()
    assert user.clerk_id is None and user.rol == 'RESIDENTE' and user.casa_unidad == 'A-001'
    assert people.query(EventOutbox).filter_by(condominio_id='a1').count() == 1
    replay = client.post(URL, headers=auth, json={**payload, 'revision': preview['revision']})
    assert replay.status_code == 409
    repeated = client.post(URL, headers=auth, json=payload).json()
    assert repeated['existentes'] == 1 and repeated['nuevos'] == 0
    assert people.query(Casa).count() == 1


@pytest.mark.parametrize('identity,role,tenant', [('guard','GUARDIA','a1'),('resident','RESIDENTE','a1'),('admin','ADMIN_CONDOMINIO','b1')])
def test_boundaries(client, people, identity, role, tenant):
    auth = headers(identity, role)
    assert client.get(f'/condominios/{tenant}/residentes/plantilla', headers=auth).status_code == 403
    assert client.post(f'/condominios/{tenant}/residentes/importar', headers=auth,
                       json=body([['1','casa','Ana','ana@example.com']])).status_code == 403


@pytest.mark.parametrize('rows', [
    [['1','casa','Ana','bad']],
    [['1','casa','Ana','ana@example.com'], ['1','casa','Luis','luis@example.com']],
    [['1','casa','Ana','ana@example.com'], ['2','casa','Luis','ANA@example.com']],
    [['1','administracion','Ana','ana@example.com']],
    [['1','casa','=1+1','ana@example.com']],
    [['1','casa','Ana','guard@test.local']],
])
def test_errors_block_whole_batch(client, people, rows):
    auth = headers('admin','ADMIN_CONDOMINIO'); payload = body(rows)
    response = client.post(URL, headers=auth, json=payload)
    assert response.status_code == 200, response.text
    preview = response.json(); assert preview['errores'] > 0
    assert client.post(URL, headers=auth,json={**payload,'revision':preview['revision']}).status_code == 400
    assert people.query(Casa).count() == 0


def test_stale_preview_and_invalid_files(client, people):
    auth = headers('admin','ADMIN_CONDOMINIO'); payload = body([['1','casa','Ana','ana@example.com']])
    preview = client.post(URL, headers=auth, json=payload).json()
    people.add(Casa(casa_id='new', condominio_id='a1', numero='1', tipo='depto')); people.commit()
    assert client.post(URL, headers=auth,json={**payload,'revision':preview['revision']}).status_code == 409
    assert client.post(URL, headers=auth,json={'archivo':'invalid'}).status_code == 400
    assert client.post(URL, headers=auth,json=body([])).status_code == 400
    assert client.get('/condominios/a1/residentes/plantilla',headers=auth).status_code == 200


def test_audit_failure_rolls_back(client, people, monkeypatch):
    import backend.routers.resident_import_router as module
    auth = headers('admin','ADMIN_CONDOMINIO'); payload = body([['001','casa','Ana','ana@example.com']])
    preview = client.post(URL, headers=auth,json=payload).json()
    from sqlalchemy.exc import IntegrityError
    def fail(*args, **kwargs): raise IntegrityError('test', {}, Exception('failure'))
    monkeypatch.setattr(module, 'encolar_evento', fail)
    assert client.post(URL,headers=auth,json={**payload,'revision':preview['revision']}).status_code == 409
    assert people.query(Casa).count() == 0
    assert people.query(Usuario).filter_by(email='ana@example.com').count() == 0
