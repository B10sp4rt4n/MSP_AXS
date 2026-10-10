from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.orm import sessionmaker

from backend.db.event import Event
from backend.services.event_outbox import enviar_pendientes
from tests.test_reglas_acceso import setup, headers
from tests.test_simulacion_reloj import reloj, crear

URL = "/reportes/a/operacion?desde=2030-01-15&hasta=2030-01-15"


def admin():
    return headers("admin", "ADMIN_CONDOMINIO")


def test_reporte_recorrido_real_y_pendientes(client, setup, reloj, db_event_engine):
    row, qr = crear(client, setup, reloj)
    assert client.get(URL, headers=admin()).json()["pendientes_entrega_condominio"] == 2
    reloj.avanzar(minutes=30)
    guard = headers("guard", "GUARDIA")
    assert client.get(qr, headers=guard).status_code == 200
    assert client.get(qr, headers=guard).status_code == 400
    assert client.patch(f"/visitas/{row.visita_id}/salida", headers=guard).status_code == 200
    other, _ = crear(client, setup, reloj)
    assert client.patch(f"/visitas/{other.visita_id}/cancelar", headers=admin()).status_code == 200
    factory = sessionmaker(bind=db_event_engine)
    enviar_pendientes(setup, "a", factory, now=reloj.utcnow(), limit=100)
    enviar_pendientes(setup, "a", factory, now=reloj.utcnow(), limit=100)
    result = client.get(URL, headers=admin())
    assert result.status_code == 200, result.text
    report = result.json()
    assert report["resumen"] == {"entrada": 1, "salida": 1, "cancelacion": 1, "rechazo": 1, "otro": 4}
    assert report["total"] == 8 and report["pendientes_entrega_condominio"] == 0
    assert report["total"] == sum(report["resumen"].values())
    assert result.headers["cache-control"] == "no-store"
    assert "qr_token" not in result.text and "session_hash" not in result.text
    assert client.get(URL + "&actor_id=guard", headers=admin()).json()["total"] == 3
    assert client.get(URL + "&destino=101&estado=salida_registrada", headers=admin()).json()["resumen"]["salida"] == 1


def event(db, uid, timestamp, **changes):
    values = dict(event_uid=uid, identity_id="guard", tenant_id="a", tipo_evento="registrar",
                  entidad="visita", entidad_id=uid, accion="registrar", resultado="exito",
                  timestamp=datetime.fromisoformat(timestamp), hash_evento=uid,
                  metadata_json={"estado": "entrada_registrada", "casa_unidad": "101"})
    db.add(Event(**{**values, **changes}))
    db.commit()


def test_corte_cdmx_filtros_y_aislamiento(client, setup, db_event_session):
    event(db_event_session, "antes", "2030-01-15T05:59:59")
    event(db_event_session, "inicio", "2030-01-15T06:00:00")
    event(db_event_session, "ultimo", "2030-01-16T05:59:59")
    event(db_event_session, "despues", "2030-01-16T06:00:00")
    event(db_event_session, "extranjero", "2030-01-15T12:00:00", tenant_id="b")
    event(db_event_session, "sin-meta", "2030-01-15T12:00:00", metadata_json=None, resultado="denegado")
    report = client.get(URL, headers=admin()).json()
    assert {e["visita_id"] for e in report["items"]} == {"inicio", "ultimo", "sin-meta"}
    assert report["resumen"]["entrada"] == 2 and report["resumen"]["rechazo"] == 1
    assert client.get(URL + "&destino=101", headers=admin()).json()["total"] == 2
    for formato in ("json", "pdf", "xlsx"):
        assert client.get(URL.replace("/a/", "/b/") + f"&formato={formato}", headers=admin()).status_code == 403
        for user, role in [("guard", "GUARDIA"), ("resident", "RESIDENTE")]:
            assert client.get(URL + f"&formato={formato}", headers=headers(user, role)).status_code == 403
    assert client.get(URL).status_code in (401, 403)


def test_validacion_y_limite_sin_truncamiento(client, setup, db_event_session, monkeypatch):
    for query in ["desde=2030-01-16&hasta=2030-01-15", "desde=2030-01-01&hasta=2030-02-01",
                  "desde=x&hasta=2030-01-15", "desde=9999-12-31&hasta=9999-12-31"]:
        assert client.get("/reportes/a/operacion?" + query, headers=admin()).status_code == 422
    assert client.get(URL + "&estado=inventado", headers=admin()).status_code == 422
    event(db_event_session, "one", "2030-01-15T12:00:00")
    event(db_event_session, "two", "2030-01-15T12:00:00")
    monkeypatch.setattr("backend.routers.reportes_router.MAX_EVENTOS", 1)
    for formato in ("json", "xlsx", "pdf"):
        assert client.get(URL + f"&formato={formato}", headers=admin()).status_code == 422


def test_exportaciones_y_texto_no_ejecutable(client, setup, db_event_session):
    event(db_event_session, "export", "2030-01-15T18:00:00", motivo="=HYPERLINK(\"https://invalid\")",
          metadata_json={"estado": "entrada_registrada", "casa_unidad": "101", "proposito": "<b>Revisión & entrega</b>"})
    event(db_event_session, "extranjero", "2030-01-15T18:00:00", tenant_id="b")
    excel = client.get(URL + "&formato=xlsx", headers=admin())
    assert excel.status_code == 200
    sheet = load_workbook(BytesIO(excel.content)).active
    cell = sheet.cell(sheet.max_row, 10)
    assert cell.value.startswith("=HYPERLINK") and cell.data_type == "s"
    assert sheet.cell(sheet.max_row, 1).value == datetime(2030, 1, 15, 12)
    assert "extranjero" not in str(list(sheet.values))
    pdf = client.get(URL + "&formato=pdf", headers=admin())
    assert pdf.status_code == 200
    reader = PdfReader(BytesIO(pdf.content))
    assert tuple(reader.pages[0].mediabox[2:]) == (612, 792)
    text = "".join(p.extract_text() for p in reader.pages)
    assert "export" in text and "extranjero" not in text and "12:00:00" in text
    assert "<b>Revisión & entrega</b>" in text


def test_vacio_exportable(client, setup):
    assert client.get(URL, headers=admin()).json()["total"] == 0
    for formato in ("xlsx", "pdf"):
        result = client.get(URL + f"&formato={formato}", headers=admin())
        assert result.status_code == 200 and len(result.content) > 500
