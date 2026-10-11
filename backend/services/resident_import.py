"""Bounded Excel parsing and read-only planning for resident imports."""
import hashlib
import io
import json
import re
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from fastapi import HTTPException
from sqlalchemy import func
from backend.db.core import Casa, Usuario, UserTenantScope, ScopeStatus, AccessLevel

HEADERS = ["vivienda", "tipo", "nombre", "correo"]
MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500


def parse_excel(raw):
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, "El archivo excede 2 MB")
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 100 or sum(i.file_size for i in archive.infolist()) > 10 * MAX_BYTES:
                raise ValueError("Archivo demasiado grande al descomprimir")
            if any("vbaProject" in i.filename for i in archive.infolist()):
                raise ValueError("No se admiten macros")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if "Residentes" not in book.sheetnames:
                raise ValueError("Falta la hoja Residentes")
            sheet = book["Residentes"]
            # Some valid exporters omit the optional worksheet dimension.
            # Discover it before enforcing the same row/column limits.
            if sheet.max_row is None or sheet.max_column is None:
                sheet.calculate_dimension(force=True)
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) != 4:
                raise ValueError("Usa cuatro columnas y máximo 500 filas de datos")
            rows = sheet.iter_rows()
            if [c.value for c in next(rows)] != HEADERS:
                raise ValueError("Encabezados requeridos: vivienda, tipo, nombre, correo")
            result = []
            for number, cells in enumerate(rows, 2):
                if all(c.value is None for c in cells):
                    continue
                errors = []
                if any(c.data_type in {"f", "e"} for c in cells):
                    errors.append("No se admiten fórmulas ni errores de Excel")
                values = [str(c.value).strip() if c.value is not None else "" for c in cells]
                unit, kind, name, email = values
                kind, email = kind.lower(), email.lower()
                if not unit or len(unit) > 100: errors.append("Vivienda requerida (máximo 100 caracteres)")
                if kind not in {"casa", "depto", "local"}: errors.append("Tipo: casa, depto o local")
                if not name or len(name) > 200: errors.append("Nombre requerido (máximo 200 caracteres)")
                if len(email) > 254 or not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email): errors.append("Correo inválido")
                result.append(dict(fila=number, vivienda=unit, tipo=kind, nombre=name, correo=email, errores=errors))
            if not result: raise ValueError("La hoja no contiene residentes")
            return result
        finally:
            book.close()
    except HTTPException:
        raise
    except (BadZipFile, ValueError, KeyError, StopIteration, OSError, TypeError) as exc:
        raise HTTPException(400, "Excel inválido: " + str(exc)) from exc
    except Exception as exc:
        raise HTTPException(400, "No se pudo leer el archivo .xlsx") from exc


def plan_import(db, tenant, rows):
    units, emails = {}, {}
    for row in rows:
        units[row['vivienda']] = units.get(row['vivienda'], 0) + 1
        emails[row['correo']] = emails.get(row['correo'], 0) + 1
    houses_by_unit = {}
    houses_all = db.query(Casa).filter(Casa.condominio_id == tenant, Casa.numero.in_(list(units))).all()
    for house in houses_all:
        houses_by_unit.setdefault(house.numero, []).append(house)
    users_by_email = {}
    users_all = db.query(Usuario).filter(func.lower(Usuario.email).in_(list(emails))).all()
    for user in users_all:
        users_by_email.setdefault(user.email.lower(), []).append(user)
    active = {s.usuario_id for s in db.query(UserTenantScope).filter(
        UserTenantScope.tenant_id == tenant, UserTenantScope.estado == ScopeStatus.ACTIVO,
        UserTenantScope.access_level == AccessLevel.RESIDENTE,
        UserTenantScope.usuario_id.in_([u.usuario_id for u in users_all])).all()}
    occupants_by_house = {}
    for occupant in db.query(Usuario).filter(Usuario.casa_id.in_([h.casa_id for h in houses_all]),
                                             Usuario.rol == "RESIDENTE").all():
        occupants_by_house.setdefault(occupant.casa_id, []).append(occupant)
    planned = []
    for source in rows:
        row = {**source, "errores": list(source['errores']), "estado": "nuevo"}
        if units[row['vivienda']] > 1: row['errores'].append("Vivienda repetida en el archivo")
        if emails[row['correo']] > 1: row['errores'].append("Correo repetido en el archivo")
        houses = houses_by_unit.get(row['vivienda'], [])
        house = houses[0] if len(houses) == 1 else None
        if len(houses) > 1: row['errores'].append("Vivienda ambigua; corrige el catálogo")
        if house and house.tipo != row['tipo']: row['errores'].append("El tipo no coincide con la vivienda existente")
        users = users_by_email.get(row['correo'], [])
        user = users[0] if len(users) == 1 else None
        if users:
            scope = user and user.usuario_id in active
            if (user and house and user.condominio_id == tenant and user.casa_id == house.casa_id
                    and user.rol == "RESIDENTE" and user.nombre == row['nombre'] and scope):
                row['estado'] = "existente"
            else:
                row['errores'].append("Correo ya registrado con otra asignación o datos; no se modificará")
        if house:
            occupants = occupants_by_house.get(house.casa_id, [])
            if any(not user or o.usuario_id != user.usuario_id for o in occupants):
                row['errores'].append("La vivienda ya tiene residente")
        row['crear_vivienda'] = not houses
        if row['errores']: row['estado'] = "error"
        planned.append(row)
    digest = hashlib.sha256(json.dumps([tenant, planned], sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return {"filas": planned, "revision": digest,
            "nuevos": sum(r['estado'] == 'nuevo' for r in planned),
            "existentes": sum(r['estado'] == 'existente' for r in planned),
            "errores": sum(r['estado'] == 'error' for r in planned)}
