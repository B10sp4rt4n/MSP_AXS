"""Exportaciones del mismo conjunto autorizado que utiliza la pantalla."""
from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

LABELS = {"entrada": "Entradas", "salida": "Salidas", "cancelacion": "Cancelaciones",
          "rechazo": "Operaciones rechazadas", "otro": "Otros eventos"}


def hora(value):
    return datetime.fromisoformat(value).astimezone(ZoneInfo("America/Mexico_City"))


def contexto(data):
    filtros = "; ".join(f"{k}: {v}" for k, v in data["filtros"].items() if v) or "Sin filtros adicionales"
    return [f'{data["condominio"]} ({data["condominio_id"]})',
            f'{data["desde"]} a {data["hasta"]} - America/Mexico_City', filtros,
            f'Generado: {hora(data["generado_en"]).strftime("%d/%m/%Y %H:%M:%S")}',
            f'Eventos pendientes de entrega del condominio: {data["pendientes_entrega_condominio"]}',
            "Hechos recibidos al generar el reporte. Los rechazos incluyen operaciones de visita y QR.",
            "Estado y destino corresponden al evento; un dato ausente no se infiere."]


def exportar(data, formato):
    output = BytesIO()
    if formato == "xlsx":
        # Runtime de la aplicación Python; no requiere servicios externos.
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "Operación"
        ws.append(["AXS - Reporte de operación"])
        for line in contexto(data):
            ws.append([line])
        for key, label in LABELS.items():
            ws.append([label, data["resumen"][key]])
        ws.append([])
        ws.append(["Fecha CDMX", "Visita", "Responsable ID", "Movimiento", "Acción", "Resultado", "Estado del evento", "Destino", "Propósito", "Motivo"])
        header = ws.max_row
        for e in data["items"]:
            ws.append([hora(e["fecha"]).replace(tzinfo=None), e["visita_id"], e["actor_id"],
                       e["categoria"], e["accion"], e["resultado"], e["estado"], e["destino"], e["proposito"], e["motivo"]])
            ws.cell(ws.max_row, 1).number_format = "dd/mm/yyyy hh:mm:ss"
        for row in ws:
            for cell in row:
                # Identificadores y texto del usuario nunca se convierten en fórmulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[header]:
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.font = Font(color="FFFFFF", bold=True)
        for row in range(1, 9):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            ws.row_dimensions[row].height = 30
        for i, width in enumerate([23, 25, 27, 18, 18, 18, 26, 22, 36, 48], 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = f"A{header+1}"
        ws.auto_filter.ref = f"A{header}:J{ws.max_row}"
        wb.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    import reportlab
    from pathlib import Path
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    fonts = Path(reportlab.__file__).parent / "fonts"
    if "AXSRegular" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("AXSRegular", str(fonts / "Vera.ttf")))
        pdfmetrics.registerFont(TTFont("AXSBold", str(fonts / "VeraBd.ttf")))
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = "AXSRegular"
    styles["Title"].fontName = styles["Heading3"].fontName = "AXSBold"
    styles["BodyText"].fontSize = 9
    styles["BodyText"].leading = 12
    def p(value):
        return Paragraph(escape(str(value if value is not None else "Sin dato")), styles["BodyText"])
    story = [Paragraph("AXS - Reporte de operación", styles["Title"])]
    story.extend(p(line) for line in contexto(data))
    story.append(Spacer(1, 12))
    summary = Table([[label, data["resumen"][key]] for key, label in LABELS.items()], colWidths=[420, 96])
    summary.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "AXSRegular"),
                                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#edf2f7")),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    story.extend([summary, Spacer(1, 16)])
    if not data["items"]:
        story.append(p("Sin eventos recibidos para este periodo y filtros."))
    for e in data["items"]:
        block = [Paragraph(f'{hora(e["fecha"]).strftime("%d/%m/%Y %H:%M:%S")} - {escape(e["categoria"])}', styles["Heading3"])]
        # Párrafos independientes permiten dividir motivos largos entre páginas.
        for label, value in [("Visita", e["visita_id"]), ("Responsable", e["actor_id"]),
                             ("Destino", e["destino"]), ("Estado del evento", e["estado"]),
                             ("Acción / resultado", f'{e["accion"]} / {e["resultado"]}'),
                             ("Propósito", e["proposito"]), ("Motivo", e["motivo"])]:
            block.append(p(f'{label}: {value if value is not None else "Sin dato"}'))
        block.append(Spacer(1, 9))
        story.append(KeepTogether(block))
    def footer(canvas, doc):
        canvas.setFont("AXSRegular", 8)
        canvas.drawString(48, 28, "AXS - Operación registrada / Horario Ciudad de México")
        canvas.drawRightString(564, 28, str(doc.page))
    SimpleDocTemplate(output, pagesize=letter, leftMargin=48, rightMargin=48,
                      topMargin=40, bottomMargin=48).build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue(), "application/pdf"
